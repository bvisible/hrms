# //// Neoffice — added file (no upstream equivalent): the text of an application's documents, for
# //// Nora's reading (neoffice-maintenance#1294).
"""The text of a document received with an application.

- **PDF and images** go through Nora's own document reader (`nora.api.ocr`): the text layer of each
  page that has one, Nora's vision for a scanned page. `document_text` never leaves Nora's primary
  engine; until it is deployed, `read_document_text` does the same job. A page the engine could not
  read raises: the reading is tried again later, never stored with a blank document.
- **Word (DOCX)** is read here, paragraphs and tables: Nora's reader takes PDF and images only.
- Without the `nora` app, a PDF's text layer is read with PyMuPDF and a scan stays "not read" —
  which the reading then says, rather than counting the document as missing.
"""

import frappe

# "Read" alone is a generic word other apps translate as a verb: the stored value says what happened
READ = "Text read"
READ_BY_OCR = "Read by OCR"
NOT_READ = "Not read"


class DocumentNotReadYet(Exception):
	"""Nora's engine could not read a page now; try the whole reading again later."""


def _path(file_url: str) -> str:
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	return file_doc.get_full_path()


def _docx_text(path: str) -> str:
	import docx

	document = docx.Document(path)
	parts = [p.text for p in document.paragraphs if p.text.strip()]
	for table in document.tables:
		for row in table.rows:
			cells = [c.text.strip() for c in row.cells if c.text.strip()]
			if cells:
				parts.append(" | ".join(cells))
	return "\n".join(parts)


def _nora_reader():
	try:
		from nora.api import ocr
	except ImportError:
		return None
	return getattr(ocr, "document_text", None) or getattr(ocr, "read_document_text", None)


def _text_layer(path: str) -> tuple[str, int]:
	import fitz

	with fitz.open(path) as pdf:
		return "\n".join(page.get_text("text") for page in pdf).strip(), pdf.page_count


def read(file_url: str) -> dict:
	"""{"text", "pages", "status"} of one document."""
	path = _path(file_url)
	extension = path.rsplit(".", 1)[-1].lower() if "." in path else ""

	if extension == "docx":
		text = _docx_text(path)
		return {"text": text, "pages": 0, "status": READ if text.strip() else NOT_READ}

	if extension not in ("pdf", "jpg", "jpeg", "png"):
		return {"text": "", "pages": 0, "status": NOT_READ}

	reader = _nora_reader()
	if reader:
		try:
			result = reader(path)
		except ValueError:
			return {"text": "", "pages": 0, "status": NOT_READ}
		except Exception as exc:
			raise DocumentNotReadYet(str(exc)[:200]) from exc
		text = (result.get("markdown") or "").strip()
		if not text:
			return {"text": "", "pages": result.get("pages") or 0, "status": NOT_READ}
		return {
			"text": text,
			"pages": result.get("pages") or 0,
			"status": READ_BY_OCR if result.get("vision_pages") else READ,
		}

	if extension == "pdf":
		text, pages = _text_layer(path)
		return {"text": text, "pages": pages, "status": READ if text else NOT_READ}
	return {"text": "", "pages": 0, "status": NOT_READ}
